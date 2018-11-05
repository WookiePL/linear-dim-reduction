import json
import os
# from pprint import pprint
# from report_model.output_report import OutputReport
import pandas as pd


class SummaryReport:
    def __init__(self):
        pass
        # self.generate_summary_report(run_id)

    # @classmethod
    # def from_dict(cls, dict):
    #     obj = cls()
    #     obj.__dict__.update(dict)
    #     return obj

    def generate_summary_report(self, run_id):

        output_report_list = []
        path = 'D:\private-projects\\linear-dim-reduction\\report_jsons'
        for filename in os.listdir(path):
            with open(path + '\\' + filename) as f:
                output_report = json.load(f)
            if output_report["run_id"] == run_id:
                output_report_list.append(output_report)
            #print(filename)
        #print(output_report_list)

        summary_report_df = pd.DataFrame(output_report_list)
        # for output_r in output_report_list:
        #     one_report_df = pd.DataFrame.from_dict(output_r, orient='index')
        #     summary_report_df.append(one_report_df)

        pass
        return summary_report_df


        #output_report: OutputReport
        # with open('D:\\private-projects\\linear-dim-reduction\\report_jsons\\2018-11-03 141735 output_report.json') as f:
        #     output_report = json.load(f)
        # #x = json.loads(data, object_hook=SummaryReport.from_dict)
        # #pprint(x)
        # pprint(output_report)

    def make_report_pretty(self, summary_report_df):

        # input_params = summary_report_df['input_params'].str
        # input_params_dict = json.load(input_params)

        input_params_list = []
        for i, r in enumerate(summary_report_df['input_params']):

            input_params_list.append(r)

        input_params_df = pd.DataFrame(input_params_list)

        pretty_report_df = summary_report_df[['accuracy', 'error', 'precision', 'recall', 'f1_score']]
        pretty_report_df.columns = ['Dokładność', 'Błąd', 'Precyzja', 'Pełność', 'Wynik F1']
        #pretty_report_df.round(2)

        pretty_report_df.insert(loc=0, column='Metoda', value=input_params_df['script_name'])
        pretty_report_df.insert(loc=1, column='Liczba_składowych', value=input_params_df['n_components'])
        pretty_report_df.insert(loc=2, column='Klasyfikator', value=input_params_df['classifier'])

        # roundings = pd.Series([2, 2, 2, 2, 2], index=['Dokładność', 'Błąd', 'Precyzja', 'Pełność', 'Wynik F1'])
        # pretty_report_df.round(roundings)
       # pd.options.display.float_format = '{:,.2f}'.format

        with open('D:\\private-projects\\linear-dim-reduction\\summary_report\\script_name_dictionary.json') as f:
            script_name_dict = json.load(f)

        for i, r in enumerate(pretty_report_df['Metoda']):
            #r1 = r.replace('derm', '_')
            old_value = pretty_report_df.at[i, 'Metoda']
            pretty_report_df.at[i, 'Metoda'] = script_name_dict[old_value]

        for i, r in enumerate(pretty_report_df['Klasyfikator']):
            #r1 = r.replace('derm', '_')
            old_value = pretty_report_df.at[i, 'Klasyfikator']
            pretty_report_df.at[i, 'Klasyfikator'] = script_name_dict[old_value]
        # pretty_report_df['Liczba składowych'].round(0)
        # pretty_report_df['Liczba składowych'] = pretty_report_df['Liczba składowych'].astype(str)

        #formatowanie kolumny z liczbą składowych
        pretty_report_df[['Liczba_składowych']] = pretty_report_df[['Liczba_składowych']].astype(str)
        for i, r in enumerate(pretty_report_df['Liczba_składowych']):
            r1 = r.replace('.0', ' ').replace('nan', '-')
            pretty_report_df.at[i, 'Liczba_składowych'] = r1

        return pretty_report_df

    def gen_docx_from_df(self, df):
        import docx
        #import pandas as pd

        # i am not sure how you are getting your data, but you said it is a
        # pandas data frame
        #df = pd.DataFrame(data)

        # open an existing document
        doc = docx.Document('D:\\private-projects\\linear-dim-reduction\\summary_report\\test2.docx')

        # add a table to the end and create a reference variable
        # extra row is so we can add the header row
        t = doc.add_table(df.shape[0]+1, df.shape[1])

        # add the header rows.
        for j in range(df.shape[-1]):
            t.cell(0,j).text = df.columns[j]

        # add the rest of the data frame
        for i in range(df.shape[0]):
            for j in range(df.shape[-1]):
                t.cell(i+1,j).text = str(df.values[i,j])

        doc.save('D:\\private-projects\\linear-dim-reduction\\summary_report\\test4.docx')


    def open_json_file(self, path):
        with open('D:\\private-projects\\linear-dim-reduction\\report_jsons\\2018-11-03 141735 output_report.json') as f:
            output_report = json.load(f)


# sr = SummaryReport()
# report_df = sr.generate_summary_report(run_id='2018-11-03 141716')

sr = SummaryReport()
report_df = sr.generate_summary_report(run_id='2018-11-05 115020')
pretty_report_df = sr.make_report_pretty(report_df)

#sr.gen_docx_from_df(report_df)



